"""Build workbook and an evidence-based draft from actual outputs; never invent metrics."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from train import compute_metrics
from eval import read_pred

SCHEMA={
 'Backbones':['exp_id','backbone','weight_tag','params_M','gmacs','img_size','epochs','seed','macro_f1_val','top1_val','train_seconds_per_epoch','latency_batch1_ms','notes'],
 'Training':['exp_id','backbone','axis','changes_from_T00','seed','macro_f1_val','top1_val','delta_T00','f1_Chinee_apple','f1_Snake_weed','notes'],
 'Inference':['exp_id','method','checkpoint','K','macro_f1','top1','ece','p50','p95','p99','images_per_s','relative_cost'],
 'Final':['exp_id','configuration','seed','macro_f1_val','macro_f1_test','top1_test','ece_test','macro_f1_test_std','top1_test_std','ece_test_std','n_seeds'],
 'PerClass':['exp_id','seed','class','test_support','precision','recall','f1'],
 'Latency':['exp_id','gpu','dtype','batch','img_size','bn_fused','p50','p95','p99','images_per_s','preprocessing','torch'],
 'Summary':['exp_id','seed','macro_f1_val','top1_val','latency_batch1_ms','notes']}


def table(df):
    if df.empty: return '*Chưa có kết quả chạy thật.*'
    columns=list(df.columns)
    def fmt(v): return f'{v:.4f}' if isinstance(v,float) else str(v).replace('|','/')
    return '\n'.join(['| '+' | '.join(columns)+' |','| '+' | '.join(['---']*len(columns))+' |']+
                     ['| '+' | '.join(fmt(v) for v in row)+' |' for row in df.itertuples(index=False,name=None)])


def generate(root):
    root=Path(root); root.mkdir(parents=True,exist_ok=True)
    rows={k:[] for k in SCHEMA}; baseline=None; configs={}
    if (root/'selected_backbone.json').exists(): baseline=json.loads((root/'selected_backbone.json').read_text())
    baseline_f1=None
    bp=root/'predictions/T00_seed0_val.csv'
    if bp.exists():
        p=read_pred(str(bp)); baseline_f1=compute_metrics(p.y_true,p.y_pred,p.probs)['macro_f1']
    for config_path in sorted((root/'runs').glob('*/seed*/config.json')):
        cfg=json.loads(config_path.read_text()); rd=config_path.parent
        summary_path=rd/'summary.json'
        if not summary_path.exists(): continue
        summary=json.loads(summary_path.read_text()); exp,seed=cfg['exp_id'],cfg['seed']
        configs[(exp,seed)]=cfg
        path=root/'predictions'/f'{exp}_seed{seed}_val.csv'
        if not path.exists(): continue
        pred=read_pred(str(path)); m=compute_metrics(pred.y_true,pred.y_pred,pred.probs)
        latency=summary['latency']; rows['Latency'].append(dict(latency,exp_id=f'{exp}_seed{seed}'))
        row=dict(exp_id=exp,seed=seed,backbone=cfg['backbone'],macro_f1_val=m['macro_f1'],top1_val=m['top1'],
                 latency_batch1_ms=latency['p50'])
        rows['Summary'].append(dict(row,notes='1-view train checkpoint; validation only'))
        if exp.startswith('B'):
            env=json.loads((rd/'environment.json').read_text())
            pre=env.get('pretrained_cfg',{})
            rows['Backbones'].append(dict(row,weight_tag=pre.get('tag',pre.get('hf_hub_id','see environment.json')),
                params_M=summary['params_M'],gmacs=summary['gmacs'],img_size=cfg['img_size'],epochs=cfg['epochs'],
                train_seconds_per_epoch=summary['train_seconds_per_epoch'],notes=summary['gmac_note']))
        if exp.startswith('T'):
            from experiments import ABLATIONS
            axis=next((a for e,a,_ in ABLATIONS if e==exp),'baseline' if exp=='T00' else 'combination')
            changes={k:v for k,v in cfg.items() if baseline and v!=baseline[k] and k not in ['exp_id','seed']}
            rows['Training'].append(dict(row,axis=axis,changes_from_T00=json.dumps(changes),
                delta_T00=m['macro_f1']-baseline_f1 if baseline_f1 is not None else None,
                f1_Chinee_apple=m['f1'][0],f1_Snake_weed=m['f1'][7],notes='Screening: one seed; no significance claim'))
    for filename,sheet in [('inference.json','Inference'),('latency.json','Latency')]:
        if (root/filename).exists(): rows[sheet].extend(json.loads((root/filename).read_text()))
    for r in rows['Inference']:
        rows['Summary'].append(dict(exp_id=r['exp_id'],seed=0,macro_f1_val=r['macro_f1'],top1_val=r['top1'],
                                    latency_batch1_ms=r['p50'],notes=r['method']))
    from dataset import CLASS_NAMES
    for exp in ['T00','F01']:
        group=[]
        for path in sorted((root/'predictions').glob(f'{exp}_seed*_test.csv')):
            p=read_pred(str(path)); m=compute_metrics(p.y_true,p.y_pred,p.probs)
            v=read_pred(str(root/'predictions'/f'{exp}_seed{p.seed}_val.csv'))
            vm=compute_metrics(v.y_true,v.y_pred,v.probs)
            cfg=configs.get((exp,p.seed),{})
            method='single'
            if exp=='F01' and (root/'final_lock.json').exists(): method=json.loads((root/'final_lock.json').read_text())['method']
            r=dict(exp_id=exp,seed=p.seed,configuration=f"{cfg.get('backbone')} / {cfg.get('loss')} / {cfg.get('aug')} / {method}",
                   macro_f1_val=vm['macro_f1'],macro_f1_test=m['macro_f1'],top1_test=m['top1'],ece_test=m['ece'])
            group.append(r); rows['Final'].append(r)
            for i,name in enumerate(CLASS_NAMES):
                rows['PerClass'].append(dict(exp_id=exp,seed=p.seed,**{'class':name},test_support=int(m['support'][i]),
                    precision=m['precision'][i],recall=m['recall'][i],f1=m['f1'][i]))
            plot_confusion(root,exp,p.seed,m['confusion'],CLASS_NAMES)
            save_errors(root,exp,p,cfg,CLASS_NAMES)
        if group:
            r=dict(exp_id=exp,seed='mean ± std (ddof=1)',n_seeds=len(group),configuration=group[0]['configuration'])
            for metric in ['macro_f1_val','macro_f1_test','top1_test','ece_test']:
                r[metric]=float(np.mean([x[metric] for x in group]))
                r[metric+'_std']=float(np.std([x[metric] for x in group],ddof=1)) if len(group)>1 else None
            rows['Final'].append(r)
    dfs={k:pd.DataFrame(v).reindex(columns=SCHEMA[k]) for k,v in rows.items()}
    dfs['Summary']=dfs['Summary'].sort_values('macro_f1_val',ascending=False).head(10)
    with pd.ExcelWriter(root/'results.xlsx',engine='openpyxl') as writer:
        for name,df in dfs.items():
            df.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='234E70')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(48,max(14,max(len(str(c.value or '')) for c in col)+2))
                for cell in col[1:]:
                    if isinstance(cell.value,float): cell.number_format='0.0000'
            metric='macro_f1_val' if 'macro_f1_val' in df else 'macro_f1' if 'macro_f1' in df else None
            if metric and not df.empty:
                index=df[metric].astype(float).idxmax()
                position=df.index.get_loc(index)+2
                for cell in ws[position]: cell.fill=PatternFill('solid',fgColor='D5F5E3')
    complete=sum(r.get('n_seeds',0)>=3 for r in rows['Final'])==2
    status='Đã có kết quả chung kết; cần đọc ảnh và bổ sung nhận xét cá nhân trước khi nộp.' if complete else 'CHƯA HOÀN TẤT THỰC NGHIỆM: các ô trống là chưa đo, không phải điểm 0.'
    sections=['# Lab Day 2 — Nguyễn Đức Triều · 2A202602978',status,
        '## 1. Tóm tắt', 'So sánh 5 backbone với cùng công thức; ablation 3 trục A/B/C; 1-view và 4 biến thể suy luận. '
        'Chọn theo macro-F1 val; khi hòa chọn ECE thấp hơn rồi độ trễ thấp hơn. Test chỉ mở sau khi khóa cấu hình. '
        'Chưa có đủ kết quả thì chưa thể kết luận cấu hình nào tốt hơn.',
        '## 2. Dữ liệu và thiết lập','DeepWeeds fold 0 nguyên bản, 9 lớp; không gộp val vào train. '
        'AdamW, 12 epoch mặc định, batch 64, LR backbone/head 1e-4/1e-3, weight decay 0.05 (trừ bias/norm), '
        'warmup 1 epoch rồi cosine theo bước. ImageNet mean/std; train crop+flip 224, val resize 256 + center crop 224. '
        'Cấu hình thực tế và phiên bản được lưu riêng trong runs/<exp_id>/seed<k>/. '
        'Macro-F1 trung bình 9 lớp, ECE 15 bin, std mẫu ddof=1. Seed sàng 0; chung kết 0,1,2. '
        'GMAC là ước lượng THOP, có thể thiếu toán tử attention; không coi đó là phép đếm chính xác cho transformer.',
        '![Phân bố lớp](eda/class_counts.png)\n\n![Ảnh theo lớp](eda/examples.png)\n\n![Augmentation](eda/augmentation.png)',
        '## 3. Backbone',table(dfs['Backbones']),
        '## 4. Công thức huấn luyện',table(dfs['Training']),
        'T01–T08 thay một yếu tố từ T00. T09 kết hợp lựa chọn val tốt nhất theo từng trục; '
        'đây là tìm kiếm cấu hình, không chứng minh hiệu ứng độc lập. Chênh lệch sàng chỉ từ một seed.',
        '## 5. Suy luận',table(dfs['Inference']),
        'Đo forward và gộp view trên thiết bị, không tính đọc ảnh, tiền xử lý CPU hay truyền CPU→GPU. '
        'Warmup 10, đo 100 lượt, đồng bộ CUDA; batch 1 và 32. Five-crop dùng ảnh resize 256, crop 224 ở bốn góc và giữa. '
        'Temperature được khớp riêng trên val mỗi seed; ECE val sau fit là số nội mẫu.',
        '## 6. Chung kết',table(dfs['Final']),
        'Ma trận nhầm lẫn và ảnh lỗi được xuất vào analysis/ theo từng seed. '
        'Không điều chỉnh mô hình sau khi xem các ảnh này.',
        '## 7. Kết luận và khuyến nghị','Chỉ kết luận sau khi có đầy đủ 3 seed cho cả mốc và cấu hình cuối. '
        'So sánh chênh lệch macro-F1 với std; nếu nhỏ hơn nhiễu, ghi “không phân biệt được”. '
        'Với robot, lọc cấu hình theo p95 đo trên thiết bị đích ≤ ngân sách 30–100 ms, rồi chọn macro-F1 val cao nhất. '
        'Thời gian trên GPU Colab không thay thế đo trên robot.',
        '## 8. Hạn chế','Một fold; sàng một seed; chọn nhiều cấu hình trên cùng val có thể quá khớp val. '
        'Chia ngẫu nhiên không bảo đảm tổng quát hóa sang địa điểm/mùa khác. '
        'Nên chạy nhiều fold và đo trực tiếp trên thiết bị triển khai. '
        'Chưa có quan sát trực tiếp của người học về các cặp lớp dễ nhầm; cần bổ sung dựa trên ảnh EDA và ảnh lỗi.',
        '## 9. Phụ lục','Notebook: [code/lab_day2.ipynb](code/lab_day2.ipynb). '
        'Các bảng trên sinh từ log và predictions bằng code/report.py; results.xlsx chứa cùng dữ liệu. '
        'eval.py gốc được dùng để score/grade; grade.txt lưu đầu ra tự chấm.']
    if (root/'eda/stats.json').exists():
        stats=json.loads((root/'eda/stats.json').read_text())
        sections.insert(7,'EDA thực tế: '+json.dumps(stats,ensure_ascii=False))
    (root/'report.md').write_text('\n\n'.join(sections),encoding='utf-8')
    if rows['Inference']:
        import matplotlib.pyplot as plt
        fig,ax=plt.subplots()
        for r in rows['Inference']:
            ax.scatter(r['p50'],r['macro_f1']); ax.annotate(r['exp_id'],(r['p50'],r['macro_f1']))
        ax.set_xlabel('Batch 1 p50 (ms)'); ax.set_ylabel('Validation macro-F1')
        fig.tight_layout(); fig.savefig(root/'inference_tradeoff.png',dpi=150); plt.close(fig)
    print(f'Wrote {root / "results.xlsx"}; {status}')


def plot_confusion(root,exp,seed,cm,names):
    import matplotlib.pyplot as plt
    out=root/'analysis'; out.mkdir(exist_ok=True)
    fig,ax=plt.subplots(figsize=(10,8)); image=ax.imshow(cm,cmap='Blues'); fig.colorbar(image,ax=ax)
    ax.set_xticks(range(9),names,rotation=65,ha='right'); ax.set_yticks(range(9),names)
    ax.set_xlabel('Predicted'); ax.set_ylabel('True'); ax.set_title(f'{exp} seed {seed} test')
    for i in range(9):
        for j in range(9): ax.text(j,i,str(cm[i,j]),ha='center',va='center',fontsize=7)
    fig.tight_layout(); fig.savefig(out/f'{exp}_seed{seed}_confusion.png',dpi=150); plt.close(fig)


def save_errors(root,exp,p,cfg,names):
    errors=np.flatnonzero(p.y_true!=p.y_pred)
    if not len(errors): return
    # read_pred retains filenames; report images only when originals are still available.
    frame=pd.read_csv(root/'predictions'/f'{exp}_seed{p.seed}_test.csv')
    frame.iloc[errors].to_csv(root/'analysis'/f'{exp}_seed{p.seed}_errors.csv',index=False)
    if not cfg or not Path(cfg['images_dir']).exists(): return
    from PIL import Image
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(3,3,figsize=(10,10))
    for ax in axes.flat: ax.axis('off')
    for ax,i in zip(axes.flat,errors[:9]):
        with Image.open(Path(cfg['images_dir'])/frame.iloc[i].Filename) as img: ax.imshow(img)
        ax.set_title(f'True: {names[p.y_true[i]]}\nPred: {names[p.y_pred[i]]}',fontsize=9)
    fig.tight_layout(); fig.savefig(root/'analysis'/f'{exp}_seed{p.seed}_errors.png',dpi=120); plt.close(fig)


if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser(); p.add_argument('--root',default='.'); generate(p.parse_args().root)
